#!/usr/bin/env python3
"""Importa la fase 'Intensità ottobre-dicembre' da "Schede mie.xlsx".

Uso:
    python3 tools/import_intensita_ottobre_dicembre.py --xlsx "<percorso>/Schede mie.xlsx"

Cosa fa (idempotente: può essere rilanciato, sostituisce le schede IOD-* esistenti):
 1. Legge il foglio 'Intensità ottobre-dicembre' (Schede A-D, 8 settimane ciascuna).
 2. Aggiorna PROGRAM_LIBRARY in src/app-main.js con 32 schede setttimanali le cui
    prescrizioni sono VERBATIM dall'Excel (serie/reps/RIR cella per cella) e con
    progression.weeks verbatim (source: "excel") su ogni esercizio.
 3. Aggiorna PROGRESSIONI-SCHEDE.json con i 4 gruppi della fase mappati ai metodi
    "Intensità: ...".

Non tocca mai i carichi: nessun default generico (6-12 / 10-20) viene scritto.
"""
import argparse
import datetime
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
PHASE = 'Intensità ottobre-dicembre'
CODE_PREFIX = 'IOD'

METHOD_NOTES = {
    'intensity-test12-wave': (
        'INTENSITÀ · TEST 12RM + ONDATA 10-X: settimana 1 test 12rm e poi 1x8 col carico '
        'trovato; poi si ondula (3×10-8-x, 2×10-x, 2×10) salendo di carico a ogni seduta, '
        'senza cambiare le reps.'
    ),
    'intensity-test10-climb': (
        'INTENSITÀ · TEST 10RM + SALITA 6→9: settimana 1 test 10rm e poi 1x6; le reps salgono '
        'di 1 ogni due settimane (6, 7, 8, 9) e il carico sale solo con tecnica impeccabile.'
    ),
    'intensity-midcycle-test': (
        'INTENSITÀ · TEST A METÀ CICLO: tre settimane di lavoro 8-10, settimana 4 test 10rm, '
        'poi si riparte più pesante con le reps che scendono (7, 8, 9).'
    ),
}
SCHEMES = {
    'intensity-test12-wave': 'Intensità: test 12rm + ondata 10-x',
    'intensity-test10-climb': 'Intensità: test 10rm + salita 6→9',
    'intensity-midcycle-test': 'Intensità: test a metà ciclo',
}
SCHEME_BY_LETTER = {
    'A': 'intensity-test12-wave',
    'B': 'intensity-test10-climb',
    'C': 'intensity-test12-wave',
    'D': 'intensity-midcycle-test',
}


def cell_str(ws, row, col):
    value = ws.cell(row=row, column=col).value
    if value is None:
        return ''
    return str(value).strip()


def parse_sheet(ws):
    blocks = {}
    current = None
    for row in range(1, ws.max_row + 1):
        marker = cell_str(ws, row, 3)
        name = cell_str(ws, row, 4)
        match = re.match(r'^Scheda\s+([A-Z])$', marker, re.IGNORECASE)
        if match:
            letter = match.group(1).upper()
            current = {'letter': letter, 'header_row': row + 1, 'exercises': []}
            blocks[letter] = current
            continue
        if current is None or row == current['header_row']:
            continue
        if not name or name.lower() == 'esercizio':
            continue
        weeks = []
        for week in range(1, 9):
            sets = cell_str(ws, row, 9 + (week - 1) * 3)
            reps = cell_str(ws, row, 10 + (week - 1) * 3)
            rir = cell_str(ws, row, 11 + (week - 1) * 3)
            if sets or reps or rir:
                weeks.append({'week': week, 'sets': sets, 'reps': reps, 'rir': rir})
        if not weeks:
            continue
        current['exercises'].append({
            'name': name,
            'muscle': cell_str(ws, row, 2),
            'note1': cell_str(ws, row, 5),
            'note2': cell_str(ws, row, 6),
            'tempo': cell_str(ws, row, 7),
            'rest': cell_str(ws, row, 8),
            'weeks': weeks,
        })
    return blocks


def week_of(exercise, week):
    for item in exercise['weeks']:
        if item['week'] == week:
            return item
    return None


def build_cards(blocks):
    cards = []
    for letter in 'ABCD':
        block = blocks[letter]
        scheme_id = SCHEME_BY_LETTER[letter]
        scheme_name = SCHEMES[scheme_id]
        method_note = METHOD_NOTES[scheme_id]
        muscles = []
        for exercise in block['exercises']:
            if exercise['muscle'] and exercise['muscle'] not in muscles:
                muscles.append(exercise['muscle'])
        focus = ', '.join(muscles)
        for week in range(1, 9):
            exercises = []
            for exercise in block['exercises']:
                current = week_of(exercise, week) or {}
                notes = ' · '.join(t for t in (exercise['note1'], exercise['note2']) if t)
                weeks_payload = [{
                    'sets': item['sets'],
                    'reps': item['reps'],
                    'rir': item['rir'],
                    'notes': notes,
                    'note': notes,
                    'week': item['week'],
                    'weekNumber': item['week'],
                    'source': 'excel',
                } for item in exercise['weeks']]
                parts = [t for t in (exercise['note1'], exercise['note2']) if t]
                exercises.append({
                    'name': exercise['name'],
                    'muscle': exercise['muscle'] or 'Custom',
                    'sets': current.get('sets', ''),
                    'reps': current.get('reps', ''),
                    'rir': current.get('rir', ''),
                    'rest': exercise['rest'],
                    'warmup': '--',
                    'tempo': exercise['tempo'],
                    'note': ' · '.join(parts),
                    'progression': {'weeks': weeks_payload, 'templateId': scheme_id, 'templateName': scheme_name},
                    'metadata': {
                        'note2': exercise['tempo'],
                        'excelNote1': exercise['note1'],
                        'excelNote2': exercise['note2'],
                    },
                    'today': '',
                    'ref': '',
                })
            cards.append({
                'code': f'{CODE_PREFIX}-{letter}{week}',
                'day': ord(letter) - 64,
                'letter': letter,
                'week': week,
                'name': f'Scheda {letter}{week}',
                'phase': PHASE,
                'focus': focus,
                'note': f'{PHASE} - Scheda {letter}, settimana {week} · Metodo: {scheme_name}',
                'exercises': exercises,
                'schemeId': scheme_id,
                'schemeName': scheme_name,
            })
    return cards


def patch_program_library(cards):
    path = ROOT / 'src' / 'app-main.js'
    src = path.read_text(encoding='utf-8')
    match = re.search(r'const PROGRAM_LIBRARY = (\[[\s\S]*?\]);', src)
    if not match:
        raise SystemExit('PROGRAM_LIBRARY non trovato in src/app-main.js')
    lib = json.loads(match.group(1))
    before = len(lib)
    lib = [p for p in lib if p.get('phase') != PHASE]
    kept = len(lib)
    lib.extend(cards)
    serialized = json.dumps(lib, ensure_ascii=False, separators=(',', ':'))
    src = src[:match.start()] + 'const PROGRAM_LIBRARY = ' + serialized + ';' + src[match.end():]
    path.write_text(src, encoding='utf-8')
    return before, kept, len(lib)


def patch_progressioni(cards):
    path = ROOT / 'PROGRESSIONI-SCHEDE.json'
    data = json.loads(path.read_text(encoding='utf-8'))
    groups = [g for g in data.get('groups', []) if g.get('phase') != PHASE]
    for letter in 'ABCD':
        card = next(c for c in cards if c['letter'] == letter)
        groups.append({
            'key': f'{PHASE}|{letter}',
            'phase': PHASE,
            'letter': letter,
            'focus': card['focus'],
            'weeks': '8',
            'schemeId': card['schemeId'],
            'schemeName': card['schemeName'],
        })
    data['groups'] = groups
    data['generatedAt'] = datetime.datetime.now(datetime.timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    return len(groups)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--xlsx', default='/home/user/uploaded_files/Schede mie.xlsx')
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.xlsx, data_only=True)
    if PHASE not in workbook.sheetnames:
        raise SystemExit(f'Foglio "{PHASE}" non trovato in {args.xlsx}')
    blocks = parse_sheet(workbook[PHASE])
    missing = [l for l in 'ABCD' if not blocks.get(l, {}).get('exercises')]
    if missing:
        raise SystemExit(f'Schede senza esercizi: {missing}')

    cards = build_cards(blocks)
    before, kept, total = patch_program_library(cards)
    group_count = patch_progressioni(cards)

    print(json.dumps({
        'phase': PHASE,
        'cards': len(cards),
        'exercisesPerSheet': {l: len(blocks[l]['exercises']) for l in 'ABCD'},
        'library': {'before': before, 'kept': kept, 'total': total},
        'progressioniGroups': group_count,
    }, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
